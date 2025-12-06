# -----------------------------------------------
# Municipal-Themed Capital Planning Presentation
# Generates: Municipal_Capital_Planning_Presentation.pptx
# Requirements: python-pptx  (pip install python-pptx)
# -----------------------------------------------

from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()

# Municipal color palette
TITLE_COLOR = RGBColor(31, 78, 121)     # Deep municipal blue (#1F4E79)
BODY_COLOR = RGBColor(51, 51, 51)       # Dark neutral (#333333)
ACCENT_COLOR = RGBColor(47, 117, 181)   # Lighter civic blue (#2F75B5)

# Slide content list
slides_content = [
    ("Why Capital Planning Matters",
     "🏘️ Community    💰 Money    🔧 Infrastructure\n\n"
     "Capital planning connects community needs, infrastructure, and finances."),
    
    ("What Is Capital Planning?",
     "Inventory → Priorities → Funding & Timing\n\n"
     "A roadmap for long-term, big-cost decisions."),
    
    ("Why This Matters to Councilors",
     "+----------------------+-----------------------+\n"
     "| Public Safety        | Fiscal Responsibility |\n"
     "| 🚑                   | 💵                    |\n"
     "+----------------------+-----------------------+\n"
     "| Transparency & Trust | Community Goals       |\n"
     "| 👥                   | 🎯                    |\n"
     "+----------------------+-----------------------+"),
    
    ("What’s in a Capital Plan?",
     "[ Timeline ]\n[ Costs & Funding ]\n[ Project Priorities ]\n[ Asset Inventory ]"),
    
    ("Benefits of Capital Planning",
     "✓ Saves money\n✓ Reduces risk\n✓ Improves services\n✓ Attracts grants\n✓ Stabilizes taxes & rates"),
    
    ("Planned vs. Emergency Repairs",
     "+---------------------------+-----------------------------+\n"
     "| Planned Repair (Green)    | Emergency Repair (Red)      |\n"
     "| 💵 Lower cost             | 💥 Higher cost              |\n"
     "| 📅 Predictable schedule   | ⚠️ Sudden disruption        |\n"
     "| 👍 Less impact            | 🚧 More damage & downtime   |\n"
     "+---------------------------+-----------------------------+"),
    
    ("Your Role in Capital Planning",
     "Set Priorities  →  Approve Funding  →  Monitor Progress\n"
     "      🎯                 💲                 📊"),
    
    ("Our Next Steps",
     "□ Update asset inventory\n□ Confirm priorities\n□ Validate costs & funding\n□ Approve multi-year plan"),
    
    ("Questions?",
     "?"),
    
    ("Final Thought",
     "“Capital planning protects people, saves money, and builds our future.”")
]

# Build slides
for title_text, body_text in slides_content:
    slide_layout = prs.slide_layouts[1]  # Title + content
    slide = prs.slides.add_slide(slide_layout)

    # Title formatting
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.size = Pt(40)
    title.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR
    title.text_frame.paragraphs[0].font.bold = True

    # Body formatting
    body = slide.placeholders[1]
    body.text = body_text
    for paragraph in body.text_frame.paragraphs:
        paragraph.font.size = Pt(26)
        paragraph.font.color.rgb = BODY_COLOR

# Save file
output_file = "C:/Users/tobi_kazeem/Downloads/Municipal_Capital_Planning_Presentation.pptx"
prs.save(output_file)

print(f"Presentation generated successfully: {output_file}")